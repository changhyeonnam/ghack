"""Fill the AI Builder Cup submission template with LedgerLens content."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import copy, os

# Submission links: fill these in before the final build.
VIDEO_URL = "TBD"  # public YouTube link to the demo video (<= 3 minutes)
APP_URL = "TBD"    # Cloud Run URL of the public demo (cache_only mode)

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = HERE + "/img"
SHOTS = HERE + "/shots"
prs = Presentation(HERE + "/template_plus2.pptx")

INK = RGBColor(0x0F, 0x3D, 0x4C)
TEAL = RGBColor(0x2A, 0x7F, 0x8E)
TINT = RGBColor(0xE6, 0xF2, 0xF4)
AMBER = RGBColor(0xE0, 0xA1, 0x00)
AMBER_T = RGBColor(0xFF, 0xF3, 0xCC)
RED = RGBColor(0xB3, 0x26, 0x1E)
GREY = RGBColor(0x5F, 0x63, 0x68)
LINE = RGBColor(0xD5, 0xDB, 0xDE)
TXT = RGBColor(0x20, 0x21, 0x24)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF7, 0xF9, 0xFA)
FONT = "Arial"

# ---------- helpers ----------

def set_title(slide, text):
    """Replace the template heading text (keeps its run formatting)."""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ""
            return sh
    return None


def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.05):
    """paras: list of dicts {t, size, bold, color, bullet, italic, align, space_after}"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for pd in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = pd.get("align", PP_ALIGN.LEFT)
        if pd.get("space_after") is not None:
            p.space_after = Pt(pd["space_after"])
        if pd.get("space_before") is not None:
            p.space_before = Pt(pd["space_before"])
        runs = pd["t"] if isinstance(pd["t"], list) else [(pd["t"], {})]
        for txt, ro in runs:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = FONT
            f.size = Pt(ro.get("size", pd.get("size", 11)))
            f.bold = ro.get("bold", pd.get("bold", False))
            f.italic = ro.get("italic", pd.get("italic", False))
            f.color.rgb = ro.get("color", pd.get("color", TXT))
        if pd.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.16))))
            pPr.set("indent", str(-int(Inches(0.16))))
            from pptx.oxml.ns import qn
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
    return box


def card(slide, x, y, w, h, fill=WHITE, line=LINE, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    from pptx.oxml.ns import qn
    st = s._element.find(qn("p:style"))
    if st is not None:
        s._element.remove(st)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def pill(slide, x, y, w, h, text, fill=INK, color=WHITE, size=8):
    s = card(slide, x, y, w, h, fill=fill, line=None, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color
    return s


def table(slide, x, y, w, h, rows, col_w, font=9.5, header_fill=INK, zebra=True, bold_first_col=False):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    t = gt.table
    # strip default style banding look by setting fills explicitly
    tblPr = t._tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(h / nrows)
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if ri == 0:
                c.fill.fore_color.rgb = header_fill
            else:
                c.fill.fore_color.rgb = PANEL if (zebra and ri % 2 == 0) else WHITE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            r.font.name = FONT; r.font.size = Pt(font)
            r.font.bold = (ri == 0) or (bold_first_col and ci == 0)
            r.font.color.rgb = WHITE if ri == 0 else TXT
    return gt


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


slides = list(prs.slides)
S = {i + 1: s for i, s in enumerate(slides)}  # template numbering

# ---------- Slide 1: Team details ----------
s = S[1]
for sh in s.shapes:
    if sh.has_text_frame and "Team Details" in sh.text_frame.text:
        ps = sh.text_frame.paragraphs
        fills = {"Team name: ": "Team name: Hey Merlion!", "Team leader name: ": "Team leader name: Changhyeon Nam (team: Changhyeon Nam, Hyunho Yang)",
                 "Problem Statement:": "Problem Statement: BFSI: Intelligent Risk, Fraud & Financial Experiences"}
        for p in ps:
            key = next((k for k in fills if p.runs and p.runs[0].text.strip() == k.strip()), None)
            if key:
                p.runs[0].text = fills[key]
                for r in p.runs[1:]:
                    r.text = ""
                p.runs[0].font.size = Pt(16)

# Redesigned slides (big numbers, cards, hero screenshots) inside the official template.
MONO = "Courier New"

def eyebrow(slide, text, y=1.22):
    tb(slide, 0.3, y, 9.4, 0.3, [{"t": text.upper(), "size": 9, "bold": True, "color": TEAL}])

def headline(slide, text, y=1.5, size=19, h=0.75):
    tb(slide, 0.3, y, 9.4, h, [{"t": text, "size": size, "bold": True, "color": INK}])

def bignum(slide, x, y, w, text, size=30, color=INK):
    tb(slide, x, y, w, size / 72 * 1.3, [{"t": text, "size": size, "bold": True, "color": color}], margin=0.0)

# ---------- Slide 2: Why this matters ----------
s = S[2]
set_title(s, "Why this matters")
eyebrow(s, "Southeast Asia's small businesses run the economy but cannot borrow")
# left hero number
tb(s, 0.3, 1.45, 4.3, 1.5, [{"t": "8.8%", "size": 96, "bold": True, "color": AMBER}], margin=0.0)
tb(s, 0.3, 3.05, 4.3, 0.7, [{"t": "of Vietnam's 914,000 SMEs have a bank loan.", "size": 17, "bold": True, "color": INK}], margin=0.0)
tb(s, 0.3, 3.75, 4.3, 0.8, [{"t": "Large enterprises: 47.1%. The same small firms employ most of the region.", "size": 12.5, "color": GREY}], margin=0.0)
# mini bar comparison
card(s, 0.32, 4.55, 4.26, 0.62, fill=PANEL, line=LINE)
card(s, 0.45, 4.64, 3.6 * 0.471, 0.18, fill=TEAL, line=None, radius=0.1)
tb(s, 0.45 + 3.6 * 0.471 + 0.05, 4.6, 1.6, 0.26, [{"t": "47.1% large firms", "size": 9, "color": TXT}], margin=0.0)
card(s, 0.45, 4.9, 3.6 * 0.088, 0.18, fill=AMBER, line=None, radius=0.1)
tb(s, 0.45 + 3.6 * 0.088 + 0.05, 4.86, 1.6, 0.26, [{"t": "8.8% SMEs", "size": 9, "bold": True, "color": TXT}], margin=0.0)
# right cards
cards2 = [("75.5%", "of Vietnamese firms cannot borrow without collateral. What they have is PDFs, screenshots and paper receipts.", False),
          ("70%", "of emerging-market MSMEs lack adequate financing. Reading their records by hand costs more than a small loan earns.", False),
          ("The opening", "VPBank, Vietcombank and Agribank are moving from collateral to cash-flow lending. Someone still has to read the cash flow. LedgerLens makes that reading fast, sourced and auditable.", True)]
y = 1.5
for big, d, acc in cards2:
    card(s, 4.9, y, 4.8, 1.12, fill=AMBER_T if acc else WHITE, line=AMBER if acc else LINE, lw=1.5 if acc else 1)
    tb(s, 5.02, y + 0.1, 1.45, 0.95, [{"t": big, "size": 24 if not acc else 13, "bold": True, "color": INK}], anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
    tb(s, 6.5, y + 0.08, 3.1, 1.0, [{"t": d, "size": 10.5, "color": TXT}], anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
    y += 1.22
tb(s, 0.3, 5.22, 9.4, 0.25, [{"t": "Sources: FiinGroup via Viet Nam News (Aug 2026) · VCCI Private Economic Report 2025 · IFC MSME Finance · VTV (May 2026) · ADB Asia SME Monitor 2025", "size": 7.5, "italic": True, "color": GREY}])

# ---------- Slide 4: Brief about the idea ----------
s = S[4]
set_title(s, "Brief about the idea")
headline(s, "LedgerLens turns a shoebox of receipts, screenshots and bank PDFs into a credit memo a bank can trust.", y=1.25, size=20, h=0.9)
mech = [("1", "Read with evidence", "Gemini reads every document. Every value carries document id, page and bounding box. Python computes the cash flow. A memo with unsourced numbers is rejected and regenerated.", False),
        ("2", "Challenge", "A second agent that never sees the first one's prompt runs 12 underwriting checks against the memo and raises severity-ranked objections with evidence.", True),
        ("3", "Human signs", "The underwriter accepts or dismisses each objection and signs. Every decision is logged and re-weights the next Challenger run. The applicant gets the reasons in their own language.", False)]
x = 0.3; cw = 2.95
for n, t, d, acc in mech:
    card(s, x, 2.3, cw, 2.35, fill=AMBER_T if acc else WHITE, line=AMBER if acc else LINE, lw=1.5 if acc else 1)
    pill(s, x + 0.18, 2.48, 0.4, 0.4, n, fill=AMBER if acc else INK, size=12)
    tb(s, x + 0.7, 2.45, cw - 0.85, 0.5, [{"t": t, "size": 16, "bold": True, "color": INK}], anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
    tb(s, x + 0.18, 3.05, cw - 0.36, 1.5, [{"t": d, "size": 11, "color": TXT}], margin=0.0)
    if n != "3":
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + cw + 0.04), Inches(3.33), Inches(0.2), Inches(0.28))
        a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xB9, 0xC6, 0xC3); a.line.fill.background(); a.shadow.inherit = False
        from pptx.oxml.ns import qn as _qn
        _st = a._element.find(_qn("p:style"))
        if _st is not None: a._element.remove(_st)
    x += cw + 0.275
tb(s, 0.3, 4.85, 9.4, 0.5, [{"t": [("Not a score. A memo. ", {"bold": True, "color": INK}), ("AI does the reading and the arguing; the underwriter keeps the decision.", {})], "size": 13, "color": GREY}])

# ---------- Slide 5: Opportunities ----------
s = S[5]
sh = set_title(s, "Opportunities")
for p in sh.text_frame.paragraphs[1:]:
    for r in p.runs:
        r.text = ""
sh.height = Inches(0.5)
headline(s, "Same documents. About 2 minutes per file. Every number defensible.", y=1.22, size=19, h=0.5)
nums = [("~2 min", "per file, end to end (measured, live Gemini)", False), ("~0.28", "USD per case on a cold run; USD 0 on a cached re-run", False), ("0 vs 62", "unsourced numbers in the memo: LedgerLens vs a single prompt", True)]
x = 0.3
for big, small, acc in nums:
    card(s, x, 1.8, 3.0, 1.35, fill=AMBER_T if acc else WHITE, line=AMBER if acc else LINE, lw=1.5 if acc else 1)
    tb(s, x + 0.18, 1.9, 2.7, 0.8, [{"t": big, "size": 38, "bold": True, "color": INK}], margin=0.0)
    tb(s, x + 0.18, 2.72, 2.7, 0.4, [{"t": small, "size": 10.5, "color": GREY}], margin=0.0)
    x += 3.2
lists = [("How it is different", ["Writer and verifier are separate agents, so the memo is never graded by its author", "Figures are computed by code and linked to their source line, not generated", "Underwriter decisions re-rank the next Challenger run"]),
         ("Who benefits", ["Small businesses with records but no credit file get assessed on cash flow, not collateral", "A declined applicant receives the reasons, in their own language", "Lenders can offer small-ticket loans that were uneconomic to underwrite"])]
x = 0.3
for t, items in lists:
    card(s, x, 3.3, 4.6, 2.05, fill=WHITE, line=LINE)
    tb(s, x + 0.18, 3.38, 4.3, 0.4, [{"t": t, "size": 13.5, "bold": True, "color": INK}], margin=0.0)
    tb(s, x + 0.18, 3.8, 4.3, 1.5, [{"t": it, "size": 10.5, "bullet": True, "space_after": 4} for it in items], margin=0.0)
    x += 4.8

# ---------- Slide 6: Features ----------
s = S[6]
set_title(s, "List of features offered by the solution")
headline(s, "Every feature hangs on one pipeline, and the pipeline ends with a person.", y=1.22, size=18, h=0.5)
feats = [("Intake", "Multimodal intake", "PDF, screenshots, photos. Classified and quality-flagged.", False),
         ("Extract", "Evidence-linked extraction", "Doc id, page, bbox on every value. PDF text layer verifies it.", False),
         ("Cash flow", "Deterministic model", "Inflow, outflow, trend, volatility, debt service. Python, not the model.", False),
         ("Analyze", "Clickable memo", "Cites catalog figures only. Click one, the source line opens.", False),
         ("Challenge", "Battery + 12 checks", "Image hash, balance chain, dates. Ranked objections with evidence.", True),
         ("Decide", "Underwriter workbench", "Accept or dismiss, sign. Audit log. Decisions re-weight the checks.", False),
         ("Explain", "Letter and paper memo", "Customer letter in the local language. Memo prints with footnoted sources.", False)]
n = len(feats); gap = 0.1; cw = (9.4 - gap * (n - 1)) / n
x = 0.3
# spine
card(s, 0.3, 1.98, 9.4, 0.04, fill=INK, line=None, radius=0.0)
for stage, t, d, acc in feats:
    pill(s, x + 0.1, 1.86, cw - 0.2, 0.28, stage, fill=AMBER if acc else INK, color=INK if acc else WHITE, size=8.5)
    card(s, x, 2.3, cw, 2.75, fill=AMBER_T if acc else WHITE, line=AMBER if acc else LINE, lw=1.5 if acc else 1)
    tb(s, x + 0.1, 2.4, cw - 0.2, 0.8, [{"t": t, "size": 11.5, "bold": True, "color": INK}], margin=0.0)
    tb(s, x + 0.1, 3.25, cw - 0.2, 1.75, [{"t": d, "size": 9.5, "color": TXT}], margin=0.0)
    x += cw + gap
tb(s, 0.3, 5.12, 9.4, 0.3, [{"t": [("Out of scope by design: ", {"bold": True, "color": INK}), ("credit scoring, loan pricing, core-banking integration. LedgerLens informs the decision; it does not make it.", {})], "size": 9.5, "color": GREY}])

# ---------- Slide 8: Wireframes -> the moment ----------
s = S[8]
set_title(s, "Wireframes / Mock diagrams of the proposed solution")
eyebrow(s, "The product in one interaction", y=1.18)
tb(s, 0.3, 1.4, 9.4, 0.5, [{"t": "Click the number. See the line.", "size": 26, "bold": True, "color": INK}], margin=0.0)
_sc = 3.4 / 460
_p = picture(s, f"{SHOTS}/v2_evidence_c.jpg", 0.3, 1.95, h=3.4); _p.line.color.rgb = LINE; _p.line.width = Pt(0.75)
marks = [(148, 72), (318, 88), (592, 118)]
for i, (px, py) in enumerate(marks):
    cx = 0.3 + px * _sc; cy = 1.95 + py * _sc
    pill(s, cx - 0.15, cy - 0.15, 0.3, 0.3, str(i + 1), fill=AMBER, color=INK, size=10)
steps8 = [("The memo cites a figure", "Every underlined value is a catalog figure with a source. Here: VND 82,500,000, credits from the applicant's own accounts."),
          ("One click opens the evidence", "The computation (sum of 3 credits whose sender matches the account holder) and all 3 source references: pages 2, 4 and 5."),
          ("The source line lights up", "The statement page opens with the extractor's bounding box drawn on the exact row. A memo with an unsourced number is rejected and regenerated.")]
y = 1.95
for i, (t, d) in enumerate(steps8):
    acc = i == 2
    card(s, 5.2, y, 4.5, 1.05, fill=AMBER_T if acc else WHITE, line=AMBER if acc else LINE, lw=1.5 if acc else 1)
    pill(s, 5.34, y + 0.14, 0.34, 0.34, str(i + 1), fill=AMBER if acc else INK, color=INK if acc else WHITE, size=11)
    tb(s, 5.8, y + 0.1, 3.8, 0.4, [{"t": t, "size": 13.5, "bold": True, "color": INK}], anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
    tb(s, 5.34, y + 0.5, 4.25, 0.55, [{"t": d, "size": 9.5, "color": TXT}], margin=0.0)
    y += 1.15

# ---------- Slide 10: Technologies ----------
s = S[10]
set_title(s, "Technologies to be used in the solution")
headline(s, "Gemini reads, ADK orchestrates, Python computes, Cloud Run serves.", y=1.22, size=18, h=0.5)
tech = [("Google Cloud", "Gemini 3.8 Flash", "Default for every agent. Measured USD 0.26-0.31 per cold case; USD 0 cached.", False),
        ("Google Cloud", "Gemini 3.1 Pro", "One env var moves Extractor, Analyst and Challenger to Pro.", False),
        ("Google Cloud", "Agent Development Kit", "Intake, 5 Extractors, Analyst, Challenger (investigator and writer), Explainer.", False),
        ("Google Cloud", "Cloud Run + Firebase", "FastAPI backend scales to zero; React workbench on Firebase Hosting.", False),
        ("Google Cloud", "Firestore + Storage", "Cases, extractions, memos, objections, decisions, audit, traces; documents.", False),
        ("Python", "Deterministic tools", "Cash-flow model, balance chain, image hash, PDF text-layer cross-check, date and entity checks.", False),
        ("Python", "LLM cache, budget cap", "Content-hash cache per model and prompt version; hard USD cap per process.", False),
        ("Integration", "REST API", "Any loan origination system can create cases, run stages and read memos.", True)]
cw, chh, gx, gy = 2.27, 1.55, 0.1, 0.12
for i, (tag, name, desc, acc) in enumerate(tech):
    r, c = divmod(i, 4)
    x = 0.3 + c * (cw + gx); y = 1.8 + r * (chh + gy)
    card(s, x, y, cw, chh, fill=AMBER_T if acc else WHITE, line=AMBER if acc else LINE, lw=1.5 if acc else 1)
    tb(s, x + 0.14, y + 0.1, cw - 0.28, 0.25, [{"t": tag.upper(), "size": 7.5, "bold": True, "color": AMBER if acc or tag == "Python" else TEAL}], margin=0.0)
    tb(s, x + 0.14, y + 0.33, cw - 0.28, 0.4, [{"t": name, "size": 11.5, "bold": True, "color": INK}], margin=0.0)
    tb(s, x + 0.14, y + 0.72, cw - 0.28, 0.8, [{"t": desc, "size": 9, "color": GREY}], margin=0.0)
tb(s, 0.3, 5.15, 9.4, 0.3, [{"t": "Output schema requires source {doc_id, page, bbox} on every value. No separate OCR engine: pages are rendered to images and read by Gemini.", "size": 9, "italic": True, "color": GREY}])

# ---------- Slide 13: Benchmark ----------
s = S[13]
set_title(s, "Prototype performance report / Benchmarking")
headline(s, "18 of 20 planted traps caught on live Gemini output; a single prompt on the same documents caught 3.", y=1.22, size=17, h=0.7)
# hero numbers
card(s, 0.3, 1.95, 4.5, 2.0, fill=AMBER_T, line=AMBER, lw=1.5)
tb(s, 0.45, 2.0, 2.2, 0.3, [{"t": "LEDGERLENS", "size": 9, "bold": True, "color": AMBER}], margin=0.0)
tb(s, 0.45, 2.25, 2.3, 1.2, [{"t": "18", "size": 72, "bold": True, "color": INK}], margin=0.0)
tb(s, 0.45, 3.4, 2.2, 0.5, [{"t": "of 20 traps caught", "size": 11, "color": TXT}, {"t": "95% Wilson CI 0.70-0.97", "size": 8.5, "color": GREY}], margin=0.0)
tb(s, 2.75, 2.0, 2.0, 0.3, [{"t": "SINGLE PROMPT", "size": 9, "bold": True, "color": GREY}], margin=0.0)
tb(s, 2.75, 2.25, 2.0, 1.2, [{"t": "3", "size": 72, "bold": True, "color": RGBColor(0xB0, 0xB8, 0xBE)}], margin=0.0)
tb(s, 2.75, 3.4, 2.0, 0.5, [{"t": "of 20, same documents", "size": 11, "color": GREY}, {"t": "and same model", "size": 8.5, "color": GREY}], margin=0.0)
card(s, 0.3, 4.1, 2.2, 1.25, fill=WHITE, line=LINE)
tb(s, 0.42, 4.15, 2.0, 0.6, [{"t": "0", "size": 30, "bold": True, "color": INK}], margin=0.0)
tb(s, 0.42, 4.75, 2.0, 0.55, [{"t": "unsourced numbers in the memo (single prompt: 62)", "size": 9.5, "color": GREY}], margin=0.0)
card(s, 2.6, 4.1, 2.2, 1.25, fill=WHITE, line=LINE)
tb(s, 2.72, 4.15, 2.0, 0.6, [{"t": "~2 min", "size": 30, "bold": True, "color": INK}], margin=0.0)
tb(s, 2.72, 4.75, 2.0, 0.55, [{"t": "per case cold, USD 0.26-0.31; cached re-run USD 0", "size": 9.5, "color": GREY}], margin=0.0)
rows = [["Metric", "Target", "Result"],
        ["Trap recall (strict)", "≥ 16 / 20", "18 / 20"],
        ["Unsourced numbers", "0", "0 (single prompt: 62)"],
        ["Objections on clean case A", "≤ 20%", "0"],
        ["Extraction row accuracy*", "≥ 97%", "100% (481 / 481)"],
        ["Evidence-link accuracy*", "≥ 95%", "100%"],
        ["Time per case (cold)", "≤ 180 s", "~2 min"],
        ["Cost per case (cold)", "≤ USD 0.50", "USD 0.26-0.31"]]
table(s, 5.0, 1.95, 4.7, 2.25, rows, [2.0, 1.0, 1.7], font=9, bold_first_col=True)
card(s, 5.0, 4.28, 4.7, 1.12, fill=PANEL, line=LINE)
tb(s, 5.1, 4.31, 4.52, 1.08, [{"t": [("How measured. ", {"bold": True}), ("Live gemini-3.8-flash runs, 2026-10-15/16: 3 synthetic businesses, 43 documents, 20 planted traps. Strict scoring: an objection counts only if it carries the trap's identifying amount or document, one objection per trap. ", {}), ("n = 3 synthetic businesses with team-planted traps, so read as indicative. Missed: dashboard gross vs payout, COD vs cash deposits. Recommendation matched on C; A gave approve_with_conditions (expected approve), B request_more_documents (expected approve_with_conditions). *Clean born-digital synthetic PDFs; real scans will be lower.", {})], "size": 7.5, "color": TXT}], margin=0.0)

# ---------- Slide 6: Process flow ----------
s = S[7]
set_title(s, "Process flow diagram")
picture(s, f"{IMG}/process_flow.png", 0.5, 1.28, h=3.8)
tb(s, 0.3, 5.12, 9.4, 0.3, [{"t": "Three lanes: the loan officer uploads, the agents read and challenge, the underwriter decides. Decisions flow back into the Challenger's priorities.", "size": 9.5, "italic": True, "color": GREY}])

# ---------- Slide 8: Architecture ----------
s = S[9]
set_title(s, "Architecture diagram of the proposed solution")
picture(s, f"{IMG}/architecture.png", 0.9, 1.28, h=3.95)


# ---------- Slide 11: Cost ----------
s = S[11]
set_title(s, "Estimated implementation cost")
tb(s, 0.3, 1.25, 9.4, 0.45, [{"t": "Measured: USD 0.26-0.31 per case on a cold Flash run, USD 0 on a cached re-run", "size": 14, "bold": True, "color": INK}])
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
cd = CategoryChartData()
cd.categories = ["LedgerLens cold run, case B (14 docs)", "LedgerLens cold run, case C (11 docs)", "Single-prompt baseline (no checks)", "LedgerLens cached re-run"]
cd.add_series("USD per case", (0.306, 0.257, 0.015, 0.0))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.3), Inches(1.75), Inches(5.7), Inches(3.5), cd)
ch = gf.chart; ch.has_legend = False; ch.has_title = True
ch.chart_title.text_frame.text = "Cost per case, USD (gemini-3.8-flash, live run 2026-10-15)"
tp = ch.chart_title.text_frame.paragraphs[0]; tp.runs[0].font.size = Pt(10.5); tp.runs[0].font.bold = True; tp.runs[0].font.color.rgb = INK; tp.runs[0].font.name = FONT
pl = ch.plots[0]; pl.has_data_labels = True; pl.gap_width = 55
pl.data_labels.number_format = '0.000'; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(10.5); pl.data_labels.font.bold = True; pl.data_labels.font.color.rgb = INK; pl.data_labels.font.name = FONT
ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = TEAL
for _i in (0, 1):
    _pt = ser.points[_i]; _pt.format.fill.solid(); _pt.format.fill.fore_color.rgb = AMBER
p2 = ser.points[2]; p2.format.fill.solid(); p2.format.fill.fore_color.rgb = RGBColor(0xB0, 0xB8, 0xBE)
ca = ch.category_axis; ca.tick_labels.font.size = Pt(9.5); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = TXT; ca.format.line.color.rgb = LINE; ca.has_major_gridlines = False
ca.reverse_order = True
va = ch.value_axis; va.visible = False; va.has_major_gridlines = False; va.maximum_scale = 0.38; va.minimum_scale = 0.0
rows = [["Assumption", "Value"],
        ["Gemini 3.8 Flash list price", "USD 0.75 / 3.75 per 1M tokens"],
        ["Measured cold run", "~2 min and USD 0.26-0.31 per case"],
        ["Cached re-run", "USD 0 (content-hash LLM cache)"],
        ["Public demo", "Replays the recorded live Gemini run, USD 0; only the customer letter calls Gemini (~USD 0.0015)"],
        ["Hosting", "Cloud Run scales to zero; Firestore and Storage in free tier"],
        ["Guardrail", "Budget-cap env var stops Gemini calls once spent"],
        ["Build", "2 engineers, 4 days, free tier and trial credits"]]
table(s, 6.3, 1.8, 3.4, 3.4, rows, [1.35, 2.05], font=8.5, bold_first_col=True)

def framed(slide, path, x, y, w=None, h=None):
    p = picture(slide, path, x, y, w=w, h=h)
    p.line.color.rgb = LINE; p.line.width = Pt(0.75)
    return p

# ---------- Slide 12: Snapshots, end to end ----------
s = S[12]
set_title(s, "Snapshots of the prototype")
shots = [("v2_newcase.jpg", "1", "New case", "Applicant details, then the raw documents. Intake classifies and quality-flags them."),
         ("v2_evidence_c.jpg", "2", "Click a number, see the line", "Evidence panel with computation and all sources; the statement opens with the bounding box highlighted."),
         ("v2_receipt_c.jpg", "3", "Challenger objections", "Critical: two receipt photos are the same image with an edited date. Evidence chips open the photo."),
         ("v2_decision_c.jpg", "4", "The underwriter signs", "Outcome, conditions and rationale, written to the audit log; accept/dismiss rates re-weight the checks.")]
for i, (f, n, t, d) in enumerate(shots):
    r, c = divmod(i, 2)
    x = 0.3 + c * 3.45; y = 1.3 + r * 2.1
    framed(s, f"{SHOTS}/{f}", x, y, h=1.95)
    pill(s, x + 0.08, y + 0.08, 0.28, 0.28, n, fill=AMBER, size=9)
cy = 1.3
for f, n, t, d in shots:
    pill(s, 7.25, cy + 0.02, 0.28, 0.28, n, fill=AMBER, size=9)
    tb(s, 7.6, cy - 0.04, 2.1, 0.4, [{"t": t, "size": 11, "bold": True, "color": INK}])
    tb(s, 7.25, cy + 0.34, 2.45, 0.72, [{"t": d, "size": 8.5, "color": GREY}])
    cy += 1.06

# ---------- Slide 14: Future development ----------
s = S[14]
set_title(s, "Additional details / Future development")
picture(s, f"{IMG}/roadmap.png", 0.3, 1.3, w=9.4)
tb(s, 0.3, 5.05, 9.4, 0.4, [{"t": "Humans decide, every figure and decision is traceable, the prototype uses watermarked synthetic data only, and production keeps documents and model endpoints in-region.", "size": 9.5, "italic": True, "color": GREY}])

# ---------- Slide 14: Links ----------
s = S[15]
sh = set_title(s, "Provide links to your:")
for p in sh.text_frame.paragraphs[1:]:
    for r in p.runs:
        r.text = ""
sh.height = Inches(0.5)
links = [("GitHub public repository", "github.com/changhyeonnam/ghack", "English README, architecture, run instructions, synthetic-data notice, presentation/"),
         ("Demo video (up to 3 minutes)", VIDEO_URL, "Problem in 20 seconds, click-to-source and Challenger scenes, benchmark near the end"),
         ("Final product", APP_URL, "Cloud Run service; each browser gets its own copy of the three demo cases (Reset demo). Scales to zero; demo cases replay the recorded live Gemini run (only the customer letter calls Gemini live), so judging costs nothing")]
y = 1.3
for t, u, d in links:
    card(s, 0.3, y, 9.4, 1.2, fill=PANEL, line=LINE)
    tb(s, 0.5, y + 0.15, 3.4, 0.5, [{"t": t, "size": 15, "bold": True, "color": INK}])
    tb(s, 0.5, y + 0.65, 3.4, 0.45, [{"t": u, "size": 12, "color": TEAL}])
    tb(s, 4.1, y + 0.15, 5.4, 0.95, [{"t": d, "size": 11.5, "color": GREY}], anchor=MSO_ANCHOR.MIDDLE)
    y += 1.33
tb(s, 0.3, 5.3, 9.4, 0.25, [{"t": "Links are inserted at submission on 18 October 2026.", "size": 10, "italic": True, "color": GREY}])

# ---------- Slide 16: Appendix, Challenger checks ----------
s = S[16]
src = [sh for sh in S[4].shapes if sh.has_text_frame and sh.text_frame.text.strip()][0]
new = copy.deepcopy(src._element)
s.shapes._spTree.append(new)
hd = s.shapes[-1]
hd.text_frame.paragraphs[0].runs[0].text = "Appendix: the Challenger's 12 fixed checks"
picture(s, f"{IMG}/checks_map.png", 0.3, 1.3, w=9.4)

# ---------- Slide 17: Appendix, more screens ----------
s = S[17]
set_title(s, "Appendix: portfolio, cash flow and audit views")
more = [("v2_portfolio.jpg", "Portfolio", "Cases in review, open objections and unsourced numbers across the book; filters by status."),
        ("v2_cashflow_c.jpg", "Cash flow", "Deterministic monthly model: inflow, net, trend, volatility, end-of-month share, repayment capacity."),
        ("v2_audit_c.jpg", "Audit timeline", "Every stage and agent call with version, model and prompt; the trail a regulator can read.")]
x = 0.3
for f, t, d in more:
    framed(s, f"{SHOTS}/{f}", x, 1.3, h=1.9)
    tb(s, x, 3.3, 3.05, 0.35, [{"t": t, "size": 12, "bold": True, "color": INK}])
    tb(s, x, 3.65, 3.05, 1.0, [{"t": d, "size": 9.5, "color": GREY}])
    x += 3.175
tb(s, 0.3, 4.8, 9.4, 0.5, [{"t": "Also in the build: dark and light themes, PDF text-layer cross-check on every extracted amount, content-hash LLM cache, cache_only demo mode, budget cap, printable paper memo with footnoted sources.", "size": 9.5, "italic": True, "color": GREY}])


# ---------- delete slide 3 (template instructions) (template instructions) ----------
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
target = ids[2]
prs.part.drop_rel(target.rId)
sldIdLst.remove(target)

out = os.path.join(os.path.dirname(HERE), "LedgerLens_AI_Builder_Cup_Submission.pptx") if os.path.basename(HERE) == "build" else HERE + "/LedgerLens_AI_Builder_Cup_Submission.pptx"
prs.save(out)
print("saved", out)
